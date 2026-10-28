import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# CONFIGURATION ET PALETTE DE COULEURS (Thème Cybersécurité)
# -----------------------------------------------------------------------------
COLOR_BG_DARK = RGBColor(15, 23, 42)      # Bleu nuit très sombre #0F172A
COLOR_CARD_BG = RGBColor(30, 41, 59)      # Bleu ardoise #1E293B
COLOR_TEXT_LIGHT = RGBColor(248, 250, 252)# Blanc cassé #F8FAFC
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# Gris clair #94A3B8
COLOR_PRIMARY = RGBColor(239, 68, 68)     # Rouge / Firebrick #EF4444
COLOR_SECONDARY = RGBColor(16, 185, 129)  # Vert émeraude #10B981
COLOR_ACCENT = RGBColor(59, 130, 246)     # Bleu clair #3B82F6

prs = Presentation()
prs.slide_width = Inches(13.333) # Format 16:9
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]

# Fonction utilitaire : Arrière-plan sombre
def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG_DARK

# Fonction utilitaire : Titre de slide standard
def add_slide_header(slide, title_text, category_text="PHASE 1 : BIG DATA & ANALYSE R"):
    # Categorie / Surtitre
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_PRIMARY
    
    # Titre principal
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_LIGHT

# -----------------------------------------------------------------------------
# SLIDE 1 : TITRE & INTRODUCTION
# -----------------------------------------------------------------------------
slide1 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide1)

# Boîte de titre principale
title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(3.5))
tf = title_box.text_frame
tf.word_wrap = True

p0 = tf.paragraphs[0]
p0.text = "PROJET CYBERSÉCURITÉ"
p0.font.size = Pt(14)
p0.font.bold = True
p0.font.color.rgb = COLOR_PRIMARY

p1 = tf.add_paragraph()
p1.text = "Analyse Big Data & Modélisation R"
p1.font.size = Pt(36)
p1.font.bold = True
p1.font.color.rgb = COLOR_TEXT_LIGHT
p1.space_before = Pt(10)

p2 = tf.add_paragraph()
p2.text = "Traitement des logs d'attaques, détection d'anomalies, Machine Learning & simulations"
p2.font.size = Pt(16)
p2.font.color.rgb = COLOR_TEXT_MUTED
p2.space_before = Pt(15)

p3 = tf.add_paragraph()
p3.text = "Auteur : Adam  |  ISEN Brest  |  Dépôt : Cyber-surveillance (R1 à R9)"
p3.font.size = Pt(13)
p3.font.color.rgb = COLOR_ACCENT
p3.space_before = Pt(40)

# -----------------------------------------------------------------------------
# SLIDE 2 : SOMMAIRE EXÉCUTIF
# -----------------------------------------------------------------------------
slide2 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide2)
add_slide_header(slide2, "Vue d'Ensemble de la Phase Big Data")

modules = [
    ("01. Fondations & EDA", "Étapes R1 à R3", "Nettoyage des logs, typage, premières visualisations et cartes de chaleur par protocole."),
    ("02. Spatio-Temporel", "Étapes R4 & R5", "Lissage par moyenne mobile sur 7 jours, saisonnalité et cartographie mondiale des hotspots."),
    ("03. Machine Learning", "Étapes R6 & R7", "Clustering K-Means, isolation des 5% d'outliers et classification prédictive via Random Forest."),
    ("04. Advanced Analytics", "Étapes R8 & R9", "Traitement NLP des payloads texte et simulation financière de risque Monte Carlo (VaR 95%).")
]

for i, (title, sub, desc) in enumerate(modules):
    left = Inches(0.8 + (i % 2) * 5.9)
    top = Inches(1.8 + (i // 2) * 2.6)
    
    # Carte fond
    shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.6), Inches(2.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD_BG
    shape.line.color.rgb = COLOR_CARD_BG
    
    tf = shape.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    
    p_sub = tf.add_paragraph()
    p_sub.text = sub
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = COLOR_ACCENT
    p_sub.space_before = Pt(4)
    
    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.size = Pt(12)
    p_desc.font.color.rgb = COLOR_TEXT_MUTED
    p_desc.space_before = Pt(10)

# -----------------------------------------------------------------------------
# FONCTION GÉNÉRIQUE POUR LES SLIDES DÉTAILLÉES (R1 à R9)
# -----------------------------------------------------------------------------
def create_detail_slide(title, subtitle, key_points, code_highlight=""):
    slide = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, title)
    
    # Colonne Gauche : Points clés
    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(7.0), Inches(5.0))
    tf_left = left_box.text_frame
    tf_left.word_wrap = True
    
    p_sub = tf_left.paragraphs[0]
    p_sub.text = subtitle
    p_sub.font.size = Pt(16)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_ACCENT
    
    for point in key_points:
        p = tf_left.add_paragraph()
        p.text = "• " + point
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_LIGHT
        p.space_before = Pt(12)
        
    # Colonne Droite : Carte résumé / extrait technique
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.8), Inches(4.3), Inches(5.0))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_CARD_BG
    
    tf_right = card.text_frame
    tf_right.word_wrap = True
    
    p_rt = tf_right.paragraphs[0]
    p_rt.text = "RÉSUMÉ TECHNIQUE"
    p_rt.font.size = Pt(14)
    p_rt.font.bold = True
    p_rt.font.color.rgb = COLOR_SECONDARY
    
    p_rc = tf_right.add_paragraph()
    p_rc.text = code_highlight
    p_rc.font.size = Pt(12)
    p_rc.font.color.rgb = COLOR_TEXT_MUTED
    p_rc.space_before = Pt(15)

# -----------------------------------------------------------------------------
# SLIDES D'ÉTAPES (R1 à R9)
# -----------------------------------------------------------------------------
create_detail_slide(
    "Étape R1 à R3 : Fondations, Nettoyage & EDA",
    "Nettoyage structuré et premières corrélations",
    [
        "Imputation et filtrage rigoureux des valeurs manquantes (NA).",
        "Normalisation des noms de colonnes via janitor::clean_names().",
        "Extraction de variables dérivées : heures, jours, niveaux de sévérité.",
        "Analyse de la répartition des protocoles (TCP, UDP, HTTP) et cartes de chaleur."
    ],
    "Livrable principal :\ncybersecurity_attacks_cleaned.csv\n\nPackages R clés :\ntidyverse, lubridate, janitor, scales."
)

create_detail_slide(
    "Étape R4 : Analyse Temporelle & Tendances",
    "Lissage de tendance et motifs hebdomadaires",
    [
        "Agrégation journalière du volume total d'incidents.",
        "Lissage par Moyenne Mobile (7 jours) avec le package zoo.",
        "Découplage des variations quotidiennes pour dégager la tendance de fond.",
        "Comparaison de l'activité par jour de la semaine (Semaine vs Week-end)."
    ],
    "Visuel produit :\ntendances_temporelles.png\n\nApport Métier :\nPermet au SOC d'anticiper les fenêtres d'attaques récurrentes."
)

create_detail_slide(
    "Étape R5 : Analyse Spatiale & Géolocalisation",
    "Cartographie des Hotspots de menaces",
    [
        "Détection automatique de la colonne géographique (Pays / IP / Localisation).",
        "Extraction du Top 10 des zones émettrices ou ciblées.",
        "Cartographie mondiale interactive ou statique (maps).",
        "Représentation des événements selon le niveau de sévérité."
    ],
    "Visuel produit :\nanalyse_spatiale.png\n\nApport Métier :\nOrientation des règles de filtrage géographique sur les pare-feu."
)

create_detail_slide(
    "Étape R6 : Clustering K-Means & Anomalies",
    "Segmentation non supervisée & isolation d'outliers",
    [
        "Normalisation centrée-réduite des variables numériques (scale).",
        "Groupement automatique en 3 clusters distincts via K-Means.",
        "Calcul de la distance euclidienne de chaque point à son centroïde.",
        "Isolation des 5% de points les plus éloignés comme 'Anomalies'."
    ],
    "Visuel produit :\nclustering_anomalies.png\n\nApport Métier :\nMise en lumière d'attaques complexes non étiquetées."
)

create_detail_slide(
    "Étape R7 : Modélisation Prédictive (Random Forest)",
    "Classification de sévérité en apprentissage supervisé",
    [
        "Partitionnement des données : 80% Entraînement / 20% Test (caret).",
        "Entraînement d'un modèle d'ensemble Random Forest (100 arbres).",
        "Évaluation de la précision (Accuracy) sur la matrice de confusion.",
        "Classement de l'importance des variables par le score de Gini."
    ],
    "Visuel produit :\nmodelisation_importance_variables.png\n\nApport Métier :\nTriage automatique de la priorité des alertes."
)

create_detail_slide(
    "Étape R8 & R9 : Advanced Analytics & Monte Carlo",
    "NLP sur Payloads & Quantification financière du risque",
    [
        "NLP (R8) : Tokenisation des payloads texte, retrait des stop-words et extraction des signatures malveillantes.",
        "Monte Carlo (R9) : 1 000 simulations probabilistes sur 30 jours.",
        "Combinaison des distributions de Poisson, Binomiale et Normale.",
        "Calcul de la valeur médiane du risque et de la Value at Risk (VaR 95%)."
    ],
    "Visuels produits :\nanalyse_nlp_payloads.png\nsimulation_monte_carlo.png\n\nApport Métier :\nTraduction du risque technique en k€ pour la direction."
)

# -----------------------------------------------------------------------------
# SLIDE FINALE : BILAN ET PROCHAINES ÉTAPES
# -----------------------------------------------------------------------------
slide_final = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide_final)
add_slide_header(slide_final, "Bilan de la Phase 1 & Roadmap")

# 2 colonnes comparatives
# Col 1
card1 = slide_final.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
card1.fill.solid()
card1.fill.fore_color.rgb = COLOR_CARD_BG
card1.line.color.rgb = COLOR_CARD_BG
tf1 = card1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "LIVRABLES COMPLÉTÉS (R1 - R9)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = COLOR_SECONDARY

items_done = [
    "Pipeline R propre, modulaire et documenté",
    "7 visuels graphiques haute définition (300 DPI)",
    "Modèle prédictif Random Forest opérationnel",
    "Moteur de simulation financière de risques Monte Carlo",
    "Code entièrement hébergé sur GitHub (Cyber-surveillance)"
]
for item in items_done:
    p = tf1.add_paragraph()
    p.text = "✓  " + item
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_LIGHT
    p.space_before = Pt(12)

# Col 2
card2 = slide_final.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.0))
card2.fill.solid()
card2.fill.fore_color.rgb = COLOR_CARD_BG
card2.line.color.rgb = COLOR_CARD_BG
tf2 = card2.text_frame
tf2.word_wrap = True

p = tf2.paragraphs[0]
p.text = "PROCHAINES ÉTAPES (PHASES 2 & 3)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

items_next = [
    "Phase 2 : Exposition des modèles via API REST (Flask/FastAPI)",
    "Phase 2 : Expliquabilité des modèles avec SHAP / LIME",
    "Phase 3 : Dashboard Live interactif en temps réel (R Shiny / Web)",
    "Phase 3 : Déploiement MLOps et automatisation des alertes"
]
for item in items_next:
    p = tf2.add_paragraph()
    p.text = "➔  " + item
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_LIGHT
    p.space_before = Pt(12)

# -----------------------------------------------------------------------------
# SAUVEGARDE DU FICHIER PPTX
# -----------------------------------------------------------------------------
output_filename = "Presentation_Cybersecurite_BigData.pptx"
prs.save(output_filename)
print(f"✅ Présentation PowerPoint générée avec succès : {os.path.abspath(output_filename)}")
